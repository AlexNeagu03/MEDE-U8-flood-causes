import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def set_run_style(run, size=20, bold=False, color=(31, 41, 55)):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)


def add_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(30, 39, 97)

    title = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.8), Inches(1.0)).text_frame
    p = title.paragraphs[0]
    run = p.add_run()
    run.text = "CP1 - U8"
    set_run_style(run, size=52, bold=True, color=(255, 255, 255))

    subtitle = slide.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.8), Inches(1.0)).text_frame
    p2 = subtitle.paragraphs[0]
    run2 = p2.add_run()
    run2.text = "Pipeline de incarcare, curatare si validare date pentru riscul de inundatii"
    set_run_style(run2, size=24, color=(202, 220, 252))

    info = slide.shapes.add_textbox(Inches(0.8), Inches(3.4), Inches(6.2), Inches(1.4)).text_frame
    info.text = "Surse: Copernicus EMS + ERA5-Land\nautor: [Nume student]\nData: octombrie 2026"
    for paragraph in info.paragraphs:
        for run in paragraph.runs:
            set_run_style(run, size=16, color=(255, 255, 255))


def add_bullet_slide(prs: Presentation, title: str, bullets: list[str], subtitle: str | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.3), Inches(12.2), Inches(0.7)).text_frame
    p = title_box.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run_style(run, size=34, bold=True, color=(30, 39, 97))

    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12.2), Inches(0.5)).text_frame
        sp = subtitle_box.paragraphs[0]
        srun = sp.add_run()
        srun.text = subtitle
        set_run_style(srun, size=16, color=(107, 114, 128))

    content = slide.shapes.add_textbox(Inches(0.9), Inches(1.7), Inches(11.8), Inches(5.3)).text_frame
    content.word_wrap = True
    content.clear()
    for index, bullet in enumerate(bullets):
        paragraph = content.paragraphs[0] if index == 0 else content.add_paragraph()
        paragraph.level = 0
        paragraph.text = bullet
        paragraph.space_after = Pt(10)
        for run in paragraph.runs:
            set_run_style(run, size=20)


def add_metrics_slide(prs: Presentation, report: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.3), Inches(12.2), Inches(0.7)).text_frame
    p = title.paragraphs[0]
    run = p.add_run()
    run.text = "Rezultate rulate pe setul inclus"
    set_run_style(run, size=34, bold=True, color=(30, 39, 97))

    subtitle = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12.2), Inches(0.5)).text_frame
    sp = subtitle.paragraphs[0]
    srun = sp.add_run()
    srun.text = "Metrici obtinute in outputs/quality_report.json"
    set_run_style(srun, size=16, color=(107, 114, 128))

    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.8), Inches(5.8), Inches(4.9))
    left.fill.solid()
    left.fill.fore_color.rgb = RGBColor(234, 242, 255)
    left.line.color.rgb = RGBColor(191, 213, 255)

    right = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(1.8), Inches(5.4), Inches(4.9))
    right.fill.solid()
    right.fill.fore_color.rgb = RGBColor(238, 249, 241)
    right.line.color.rgb = RGBColor(190, 231, 201)

    left_text = slide.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(5.1), Inches(4.1)).text_frame
    input_ems = report.get("input_rows", {}).get("ems_events", "-")
    input_era5 = report.get("input_rows", {}).get("era5_daily", "-")
    dup_ems = report.get("duplicates_removed", {}).get("ems_events", "-")
    dup_era5 = report.get("duplicates_removed", {}).get("era5_daily", "-")
    filtered = report.get("window_filter", {}).get("era5_removed_outside_event_window", "-")
    left_text.text = (
        "Input\n\n"
        f"EMS rows: {input_ems}\n"
        f"ERA5 rows: {input_era5}\n\n"
        "Curatare\n"
        f"Duplicate eliminate: EMS {dup_ems}, ERA5 {dup_era5}\n"
        f"Filtrare fereastra: ERA5 eliminate {filtered}"
    )
    for paragraph in left_text.paragraphs:
        for run in paragraph.runs:
            set_run_style(run, size=18)

    right_text = slide.shapes.add_textbox(Inches(7.3), Inches(2.1), Inches(4.8), Inches(4.1)).text_frame
    out_ems = report.get("final_rows", {}).get("ems_events_clean", "-")
    out_era5 = report.get("final_rows", {}).get("era5_daily_clean", "-")
    out_join = report.get("final_rows", {}).get("joined_event_daily", "-")
    right_text.text = (
        "Output final\n\n"
        f"ems_events_clean: {out_ems}\n"
        f"era5_daily_clean: {out_era5}\n"
        f"joined_event_daily: {out_join}\n\n"
        "Null counts: vezi quality_report.json"
    )
    for paragraph in right_text.paragraphs:
        for run in paragraph.runs:
            set_run_style(run, size=18)


def add_final_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(30, 39, 97)

    title = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.8), Inches(0.9)).text_frame
    p = title.paragraphs[0]
    run = p.add_run()
    run.text = "Concluzii CP1"
    set_run_style(run, size=44, bold=True, color=(255, 255, 255))

    body = slide.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11.8), Inches(3.9)).text_frame
    body.text = (
        "Am livrat un pipeline complet pentru ingestia si validarea datelor U8,\n"
        "cu artefacte gata de predare:\n"
        "- cod sursa modular\n"
        "- teste automate\n"
        "- datacard + SOA\n"
        "- fisiere procesate + raport de calitate\n\n"
        "Pas urmator (CP2): analiza exploratorie si modelare pe joined_event_daily.csv"
    )
    for paragraph in body.paragraphs:
        for run in paragraph.runs:
            set_run_style(run, size=22, color=(202, 220, 252))

    thanks = slide.shapes.add_textbox(Inches(0.9), Inches(6.2), Inches(3.5), Inches(0.5)).text_frame
    tp = thanks.paragraphs[0]
    tr = tp.add_run()
    tr.text = "Multumesc!"
    set_run_style(tr, size=24, bold=True, color=(255, 255, 255))
    tp.alignment = PP_ALIGN.LEFT


def read_report(project_root: Path) -> dict:
    report_path = project_root / "outputs" / "quality_report.json"
    if report_path.exists():
        return json.loads(report_path.read_text(encoding="utf-8"))
    return {}


def build_presentation(output_path: Path) -> None:
    project_root = output_path.resolve().parent
    report = read_report(project_root)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)
    add_bullet_slide(
        prs,
        "Context si obiective",
        [
            "Obiectiv principal: pipeline reproductibil care aduce date brute in tabele curate.",
            "Implementare: incarcare CSV pentru doua surse independente (EMS + ERA5).",
            "Validari: schema, tipuri, intervale fizice si coerenta temporala.",
            "Post-procesare: deduplicare, filtrare dupa fereastra evenimentului si join final.",
            "Rezultat: seturi curate + quality_report.json pentru trasabilitate.",
        ],
        "Ce face CP1 fata de CP0",
    )
    add_bullet_slide(
        prs,
        "Surse de date utilizate",
        [
            "S1 - Copernicus EMS: metadate eveniment (event_id, zona, start/end, hazard_type).",
            "S2 - ERA5-Land: indicatori meteo zilnici (precipitatii, temperatura, umiditate sol).",
            "Cheie comuna: event_id.",
            "Granularitate finala: 1 eveniment x 1 zi x 1 celula spatiala.",
        ],
    )
    add_bullet_slide(
        prs,
        "Structura datelor (SOA)",
        [
            "EMS (event): PK event_id; atribute descriptive + interval temporal.",
            "ERA5 (daily): PK compus (event_id, date, grid_id); atribute numerice meteo.",
            "Join final: event_id + constrangere date in [start_date, end_date].",
            "Output: joined_event_daily.csv pentru analiza CP2.",
        ],
    )
    add_bullet_slide(
        prs,
        "Arhitectura pipeline",
        [
            "Pas 1: Read CSV (data/raw).",
            "Pas 2: Validate (required fields, parse date/numeric, range checks).",
            "Pas 3: Deduplicate (event_id; event_id+date+grid_id).",
            "Pas 4: Filter temporal (ERA5 doar in fereastra evenimentului).",
            "Pas 5: Join + export in data/processed si outputs/quality_report.json.",
        ],
    )
    add_bullet_slide(
        prs,
        "Reguli de calitate implementate",
        [
            "Completitudine: campuri obligatorii nenule.",
            "Corectitudine: date ISO si valori numerice valide.",
            "Intervale fizice: limite pentru precip, temperatura, umiditate, lat/lon.",
            "Integritate temporala: date ERA5 doar in interval valid per eveniment.",
            "Trasabilitate: metrici de intrare/iesire si lista erori in raport JSON.",
        ],
    )
    add_metrics_slide(prs, report)
    add_final_slide(prs)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


if __name__ == "__main__":
    build_presentation(Path("CP1_U8_prezentare.pptx"))
